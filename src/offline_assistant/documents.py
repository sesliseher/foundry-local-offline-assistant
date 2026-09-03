"""TXT, metin tabanlı PDF ve DOCX okuma; kaynak bilgili parçalama."""

from dataclasses import dataclass
from pathlib import Path
import re


@dataclass(frozen=True)
class Document:
    source: str
    text: str
    file_type: str = "txt"
    page_number: int | None = None


@dataclass(frozen=True)
class Chunk:
    source: str
    chunk_number: int
    text: str
    file_type: str = "txt"
    page_number: int | None = None


SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}


def read_txt_documents(directory: Path, encoding: str = "utf-8-sig") -> list[Document]:
    """Alt klasörleri tarar; aynı adlı dosyaları göreli yollarıyla ayırt eder."""
    directory = directory.resolve()
    if not directory.is_dir():
        raise ValueError(f"Belge klasörü bulunamadı: {directory}")
    paths = sorted(
        (path for path in directory.rglob("*") if path.is_file() and path.suffix.lower() == ".txt"),
        key=lambda path: path.relative_to(directory).as_posix(),
    )
    documents = []
    for path in paths:
        source = path.relative_to(directory).as_posix()
        try:
            text = path.read_text(encoding=encoding)
        except UnicodeError as exc:
            raise ValueError(
                f"{source} dosyası {encoding} ile okunamadı. Dosyayı UTF-8 kaydedin "
                "veya doğru --encoding seçeneğini kullanın."
            ) from exc
        documents.append(Document(source=source, text=text))
    return documents


def read_documents(directory: Path, encoding: str = "utf-8-sig") -> list[Document]:
    """Desteklenen belgeleri okur; PDF sayfalarını ayrı metadata ile korur."""
    directory = directory.resolve()
    if not directory.is_dir():
        raise ValueError(f"Belge klasörü bulunamadı: {directory}")
    paths = sorted(
        (path for path in directory.rglob("*") if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS),
        key=lambda path: path.relative_to(directory).as_posix(),
    )
    documents = []
    for path in paths:
        source = path.relative_to(directory).as_posix()
        suffix = path.suffix.lower()
        try:
            if suffix == ".txt":
                documents.append(Document(source, path.read_text(encoding=encoding), "txt"))
            elif suffix == ".docx":
                from docx import Document as WordDocument
                word = WordDocument(path)
                text = "\n\n".join(p.text for p in word.paragraphs if p.text.strip())
                documents.append(Document(source, text, "docx"))
            else:
                from pypdf import PdfReader
                reader = PdfReader(path)
                for page_number, page in enumerate(reader.pages, start=1):
                    documents.append(Document(source, page.extract_text() or "", "pdf", page_number))
        except UnicodeError as exc:
            raise ValueError(f"{source} dosyası {encoding} ile okunamadı.") from exc
        except Exception as exc:
            raise ValueError(f"{source} belgesi okunamadı: {exc}") from exc
    return documents


def split_document(document: Document, max_chars: int = 600) -> list[Chunk]:
    """Paragrafları mümkün olduğunca korur; sınır karakter sayısıdır, token değil."""
    if max_chars < 1:
        raise ValueError("Parça boyutu sıfırdan büyük olmalıdır.")
    text = document.text.replace("\r\n", "\n").replace("\r", "\n")
    paragraphs = [
        " ".join(paragraph.split())
        for paragraph in re.split(r"\n\s*\n", text)
        if paragraph.strip()
    ]
    pieces = []
    for paragraph in paragraphs:
        # Uzun paragraflarda kelime sınırını tercih et; çok uzun tek kelimeyi böl.
        while len(paragraph) > max_chars:
            boundary = paragraph.rfind(" ", 0, max_chars + 1)
            if boundary <= 0:
                boundary = max_chars
            pieces.append(paragraph[:boundary])
            paragraph = paragraph[boundary:].lstrip()
        if paragraph:
            pieces.append(paragraph)

    texts = []
    current = ""
    for piece in pieces:
        combined = f"{current}\n\n{piece}" if current else piece
        if len(combined) <= max_chars:
            current = combined
        else:
            texts.append(current)
            current = piece
    if current:
        texts.append(current)
    return [
        Chunk(document.source, number, content, document.file_type, document.page_number)
        for number, content in enumerate(texts, start=1)
    ]


def split_documents(documents: list[Document], max_chars: int = 600) -> list[Chunk]:
    """Sayfalara ayrılmış aynı kaynağa benzersiz, artan parça numarası verir."""
    counters: dict[str, int] = {}
    chunks = []
    for document in documents:
        for chunk in split_document(document, max_chars):
            number = counters.get(chunk.source, 0) + 1
            counters[chunk.source] = number
            chunks.append(Chunk(chunk.source, number, chunk.text, chunk.file_type, chunk.page_number))
    return chunks
