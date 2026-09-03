"""TXT okuma, kaynak takibi ve metin kaybı olmadan parçalama kontrolleri."""

import pytest

from docx import Document as WordDocument

from src.offline_assistant.documents import (
    Document, read_documents, read_txt_documents, split_document, split_documents,
)


def test_paragraphs_and_source_numbers_are_preserved():
    document = Document("ders/not.txt", "Birinci paragraf.\n\nİkinci paragraf.")
    chunks = split_document(document, max_chars=20)
    assert [chunk.text for chunk in chunks] == ["Birinci paragraf.", "İkinci paragraf."]
    assert [chunk.chunk_number for chunk in chunks] == [1, 2]
    assert all(chunk.source == "ders/not.txt" for chunk in chunks)


def test_long_paragraph_keeps_all_words_in_order():
    text = "Kütüphane hafta içi açık kalır ve öğrenciler burada sessizce çalışabilir."
    chunks = split_document(Document("not.txt", text), max_chars=23)
    assert " ".join(chunk.text for chunk in chunks).split() == text.split()
    assert all(0 < len(chunk.text) <= 23 for chunk in chunks)


def test_long_single_word_is_not_lost():
    text = "ğ" * 31
    chunks = split_document(Document("not.txt", text), max_chars=10)
    assert "".join(chunk.text for chunk in chunks) == text
    assert [len(chunk.text) for chunk in chunks] == [10, 10, 10, 1]


def test_short_paragraphs_combine_and_windows_newlines_normalize():
    chunks = split_document(Document("not.txt", " Başlık\r\n\r\nAçıklama\r\nsatırı. "), 100)
    assert len(chunks) == 1
    assert chunks[0].text == "Başlık\n\nAçıklama satırı."


def test_empty_document_has_no_chunks():
    assert split_document(Document("bos.txt", " \n\t\r\n")) == []


@pytest.mark.parametrize("max_chars", [0, -1])
def test_invalid_chunk_size(max_chars):
    with pytest.raises(ValueError):
        split_document(Document("not.txt", "metin"), max_chars)


def test_reader_handles_bom_nested_names_and_uppercase_extension(tmp_path):
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    (tmp_path / "a" / "not.txt").write_text("Türkçe içerik", encoding="utf-8-sig")
    (tmp_path / "b" / "not.TXT").write_text("İkinci belge", encoding="utf-8")
    (tmp_path / "ignore.pdf").write_text("TXT değildir", encoding="utf-8")
    documents = read_txt_documents(tmp_path)
    assert [doc.source for doc in documents] == ["a/not.txt", "b/not.TXT"]
    assert [doc.text for doc in documents] == ["Türkçe içerik", "İkinci belge"]


def test_invalid_encoding_does_not_silently_drop_characters(tmp_path):
    (tmp_path / "not.txt").write_bytes("Öğrenci".encode("cp1254"))
    with pytest.raises(ValueError, match="not.txt"):
        read_txt_documents(tmp_path)
    assert read_txt_documents(tmp_path, encoding="cp1254")[0].text == "Öğrenci"


def test_missing_directory_is_an_error(tmp_path):
    with pytest.raises(ValueError, match="klasörü bulunamadı"):
        read_txt_documents(tmp_path / "missing")


def test_docx_reader_preserves_type_and_text(tmp_path):
    word = WordDocument()
    word.add_paragraph("Başlık")
    word.add_paragraph("Türkçe DOCX içeriği")
    word.save(tmp_path / "bilgi.docx")

    documents = read_documents(tmp_path)

    assert [(doc.source, doc.file_type) for doc in documents] == [("bilgi.docx", "docx")]
    assert documents[0].text == "Başlık\n\nTürkçe DOCX içeriği"


def test_split_documents_numbers_pdf_pages_without_collisions():
    documents = [
        Document("rehber.pdf", "Birinci sayfa", "pdf", 1),
        Document("rehber.pdf", "İkinci sayfa", "pdf", 2),
    ]
    chunks = split_documents(documents, 600)
    assert [(c.chunk_number, c.page_number, c.file_type) for c in chunks] == [
        (1, 1, "pdf"), (2, 2, "pdf")
    ]
